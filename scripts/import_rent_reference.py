#!/usr/bin/env python3
"""Import the authorized official e-Stat rent reference XLSX files."""

from __future__ import annotations

import argparse
import asyncio
import os
import re
import urllib.request
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import asyncpg
from openpyxl import load_workbook

from backend.app.region_names import RegionMappingReport, map_region_names

HOUSING_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040210062&fileKind=0"
HOUSING_122_5_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040210063&fileKind=0"
KOURI_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040506178&fileKind=0"
HOUSING_SOURCE = "estat_housing_land_122_4"
HOUSING_122_5_SOURCE = "estat_housing_land_122_5"
KOURI_SOURCE = "estat_kouri_3001"
LICENSE_LABEL = "e-Stat利用規約（出典明記・加工して作成）"
HOUSING_LABEL = "令和5年住宅・土地統計調査 第122-4表"
HOUSING_122_5_LABEL = "令和5年住宅・土地統計調査 第122-5表"
SCOPE_HOUSING = "民営借家・借家(専用住宅)"
SCOPE_KOURI = "民営家賃・都市別・1か月･3.3m2"
PREFECTURE_BY_CODE = {
    "01": "北海道", "02": "青森県", "03": "岩手県", "04": "宮城県", "05": "秋田県", "06": "山形県",
    "07": "福島県", "08": "茨城県", "09": "栃木県", "10": "群馬県", "11": "埼玉県", "12": "千葉県",
    "13": "東京都", "14": "神奈川県", "15": "新潟県", "16": "富山県", "17": "石川県", "18": "福井県",
    "19": "山梨県", "20": "長野県", "21": "岐阜県", "22": "静岡県", "23": "愛知県", "24": "三重県",
    "25": "滋賀県", "26": "京都府", "27": "大阪府", "28": "兵庫県", "29": "奈良県", "30": "和歌山県",
    "31": "鳥取県", "32": "島根県", "33": "岡山県", "34": "広島県", "35": "山口県", "36": "徳島県",
    "37": "香川県", "38": "愛媛県", "39": "高知県", "40": "福岡県", "41": "佐賀県", "42": "長崎県",
    "43": "熊本県", "44": "大分県", "45": "宮崎県", "46": "鹿児島県", "47": "沖縄県",
}


@dataclass
class ParseReport:
    skipped: int = 0
    unmapped_prefecture: int = 0
    unmapped_city: int = 0
    samples: list[str] = field(default_factory=list)


def _number(value: Any) -> float | None:
    if value is None or str(value).strip() in {"", "-", "…", "..."}:
        return None
    try:
        number = float(str(value).replace(",", "").strip())
    except ValueError:
        return None
    return number if number > 0 else None


def _month(value: Any) -> str | None:
    match = re.fullmatch(r"(\d{4})年(\d{1,2})月", str(value or "").strip())
    return f"{match.group(1)}-{int(match.group(2)):02d}" if match else None


def _region(source_prefecture: str, source_city: str, report: ParseReport, *, special: bool = False) -> tuple[str, str, str, str] | None:
    if special and source_prefecture == "東京都" and source_city in {"特別区部", "東京都区部"}:
        return "东京都", "东京23区", "__not_subdivided__", "special_wards"
    prefecture, city, ward = map_region_names(source_prefecture, source_city)
    if not prefecture:
        report.unmapped_prefecture += 1
        return None
    if not city:
        report.unmapped_city += 1
        if source_city not in report.samples and len(report.samples) < 10:
            report.samples.append(source_city)
        return None
    ward = ward or "__not_subdivided__"
    level = "ward" if ward != "__not_subdivided__" else "city"
    return prefecture, city, ward, level


def parse_housing_land_workbook(path: Path, fetched_at: datetime | None = None) -> tuple[list[dict[str, Any]], ParseReport]:
    report = ParseReport()
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["e122_4"]
    prefectures = {str(row[1]).split("_", 1)[0]: str(row[1]).split("_", 1)[1] for row in sheet.iter_rows(min_row=10, values_only=True) if row[0] == "a" and row[1] and str(row[1]).split("_", 1)[0] != "00000"}
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=10, values_only=True):
        if str(row[2] or "").strip() != "3_民営借家":
            continue
        incl, excl = _number(row[15]), _number(row[16])
        if incl is None or excl is None:
            report.skipped += 1
            continue
        code, source_city = str(row[1]).split("_", 1) if row[1] and "_" in str(row[1]) else ("", "")
        source_prefecture = prefectures.get(code[:2] + "000")
        if not source_prefecture:
            report.skipped += 1
            continue
        if code == code[:2] + "000":
            prefecture, city, ward, geo_level = map_region_names(source_prefecture, source_prefecture)[0], "__not_subdivided__", "__not_subdivided__", "prefecture"
        else:
            region = _region(source_prefecture, source_city.replace("　", ""), report, special=True)
            if not region:
                report.skipped += 1
                continue
            prefecture, city, ward, geo_level = region
        rows.append({
            "source_key": HOUSING_SOURCE, "source_label": HOUSING_LABEL, "prefecture": prefecture, "city": city, "ward": ward,
            "geo_level": geo_level, "scope_label": SCOPE_HOUSING, "rent_jpy_per_sqm_month": incl,
            "rent_jpy_per_sqm_month_excl_zero": excl, "survey_year": 2023, "survey_label": "令和5年(2023)",
            "observed_month": None, "source_url": HOUSING_URL, "license_label": LICENSE_LABEL,
            "fetched_at": fetched_at or datetime.now(timezone.utc),
            "building_type": None, "structure_type": None,
        })
    return rows, report


def _dimension_label(value: Any) -> str | None:
    text = str(value or "").strip()
    return text.split("_", 1)[1] if "_" in text else None


def parse_housing_land_122_5_workbook(path: Path, fetched_at: datetime | None = None) -> tuple[list[dict[str, Any]], ParseReport]:
    """Parse the measured e-Stat 122-5 layout using its header labels."""
    report = ParseReport()
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["e122_5"]
    header_rows = list(sheet.iter_rows(min_row=1, max_row=9, values_only=True))
    location_row = next((row for row in header_rows if row[:4] == ("地域識別コード", "地域区分", "住宅の建て方", "建物の構造")), None)
    if location_row is None:
        raise ValueError("122-5 header row not found")
    building_index = location_row.index("住宅の建て方")
    structure_index = location_row.index("建物の構造")
    rent_index = next(
        index for row in header_rows for index, value in enumerate(row)
        if value == "2_家賃０円を含まない"
    )
    prefectures = {
        str(row[1]).split("_", 1)[0]: str(row[1]).split("_", 1)[1]
        for row in sheet.iter_rows(min_row=10, values_only=True)
        if row[0] == "a" and row[1] and "_" in str(row[1])
        and str(row[1]).split("_", 1)[0] != "00000"
        and len(str(row[1]).split("_", 1)[0]) == 5
        and str(row[1]).split("_", 1)[0].endswith("000")
    }
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=10, values_only=True):
        if len(row) <= rent_index or not row[1] or "_" not in str(row[1]):
            report.skipped += 1
            continue
        building_type = _dimension_label(row[building_index])
        structure_type = _dimension_label(row[structure_index])
        rent = _number(row[rent_index])
        code, source_region = str(row[1]).split("_", 1)
        if not building_type or not structure_type or rent is None:
            report.skipped += 1
            continue
        source_prefecture = prefectures.get(code[:2] + "000")
        if not source_prefecture:
            report.skipped += 1
            continue
        if code == code[:2] + "000":
            prefecture = map_region_names(source_prefecture, source_prefecture)[0]
            if not prefecture:
                report.unmapped_prefecture += 1
                report.skipped += 1
                continue
            city, ward, geo_level = "__not_subdivided__", "__not_subdivided__", "prefecture"
        else:
            region = _region(source_prefecture, source_region.replace("　", ""), report, special=True)
            if not region:
                report.skipped += 1
                continue
            prefecture, city, ward, geo_level = region
        rows.append({
            "source_key": HOUSING_122_5_SOURCE, "source_label": HOUSING_122_5_LABEL,
            "prefecture": prefecture, "city": city, "ward": ward, "geo_level": geo_level,
            "scope_label": f"借家(専用住宅)・{building_type}・{structure_type}",
            "building_type": building_type, "structure_type": structure_type,
            "rent_jpy_per_sqm_month": rent, "rent_jpy_per_sqm_month_excl_zero": rent,
            "survey_year": 2023, "survey_label": "令和5年(2023)", "observed_month": None,
            "source_url": HOUSING_122_5_URL, "license_label": LICENSE_LABEL,
            "fetched_at": fetched_at or datetime.now(timezone.utc),
        })
    return rows, report


def parse_kouri_workbook(path: Path, fetched_at: datetime | None = None) -> tuple[list[dict[str, Any]], ParseReport]:
    report = ParseReport()
    rows = list(load_workbook(path, read_only=True, data_only=True)["B001-2"].iter_rows(values_only=True))
    codes, cities = rows[9], rows[10]
    parsed: list[dict[str, Any]] = []
    for row in rows:
        if row[9] != 3001:
            continue
        observed_month = _month(row[8])
        for index in range(15, min(96, len(row))):
            value = _number(row[index])
            if value is None or not codes[index] or not cities[index] or not observed_month:
                report.skipped += 1
                continue
            code, source_city = str(codes[index]), str(cities[index]).strip()
            source_prefecture = PREFECTURE_BY_CODE.get(code[:2])
            region = _region(source_prefecture or "", source_city, report, special=True) if source_prefecture else None
            if not region:
                report.skipped += 1
                continue
            prefecture, city, ward, geo_level = region
            rent = value / 3.3
            parsed.append({
                "source_key": KOURI_SOURCE, "source_label": f"小売物価統計調査(動向編) 2026年8月", "prefecture": prefecture,
                "city": city, "ward": ward, "geo_level": geo_level, "scope_label": SCOPE_KOURI,
                "rent_jpy_per_sqm_month": rent, "rent_jpy_per_sqm_month_excl_zero": rent, "survey_year": 2026,
                "survey_label": "2026年8月", "observed_month": observed_month, "source_url": KOURI_URL,
            "license_label": LICENSE_LABEL, "fetched_at": fetched_at or datetime.now(timezone.utc),
                "building_type": None, "structure_type": None,
            })
    return parsed, report


IMPORT_COLUMNS = ("source_key", "source_label", "prefecture", "city", "ward", "geo_level", "scope_label", "building_type", "structure_type", "rent_jpy_per_sqm_month", "rent_jpy_per_sqm_month_excl_zero", "survey_year", "survey_label", "observed_month", "source_url", "license_label", "fetched_at")


async def upsert_rows(database_url: str, rows: list[dict[str, Any]]) -> dict[str, int]:
    conn = await asyncpg.connect(database_url)
    try:
        await conn.execute("create temporary table _rent_reference_import (like public.rent_reference_stats including defaults) on commit preserve rows")
        await conn.copy_records_to_table("_rent_reference_import", records=(tuple(row[column] for column in IMPORT_COLUMNS) for row in rows), columns=IMPORT_COLUMNS)
        updated = await conn.fetchval("""with changed as (update public.rent_reference_stats target set source_label=source.source_label, building_type=source.building_type, structure_type=source.structure_type, rent_jpy_per_sqm_month=source.rent_jpy_per_sqm_month, rent_jpy_per_sqm_month_excl_zero=source.rent_jpy_per_sqm_month_excl_zero, survey_label=source.survey_label, source_url=source.source_url, license_label=source.license_label, fetched_at=source.fetched_at from _rent_reference_import source where target.source_key=source.source_key and target.prefecture=source.prefecture and target.city=source.city and target.ward=source.ward and target.scope_label=source.scope_label and target.building_type is not distinct from source.building_type and target.structure_type is not distinct from source.structure_type and target.survey_year=source.survey_year and target.observed_month is not distinct from source.observed_month and (target.source_label, target.building_type, target.structure_type, target.rent_jpy_per_sqm_month, target.rent_jpy_per_sqm_month_excl_zero, target.survey_label, target.source_url, target.license_label) is distinct from (source.source_label, source.building_type, source.structure_type, source.rent_jpy_per_sqm_month, source.rent_jpy_per_sqm_month_excl_zero, source.survey_label, source.source_url, source.license_label) returning 1) select count(*) from changed""")
        inserted = await conn.fetchval("""with added as (insert into public.rent_reference_stats (source_key,source_label,prefecture,city,ward,geo_level,scope_label,building_type,structure_type,rent_jpy_per_sqm_month,rent_jpy_per_sqm_month_excl_zero,survey_year,survey_label,observed_month,source_url,license_label,fetched_at) select source_key,source_label,prefecture,city,ward,geo_level,scope_label,building_type,structure_type,rent_jpy_per_sqm_month,rent_jpy_per_sqm_month_excl_zero,survey_year,survey_label,observed_month,source_url,license_label,fetched_at from _rent_reference_import source where not exists (select 1 from public.rent_reference_stats target where target.source_key=source.source_key and target.prefecture=source.prefecture and target.city=source.city and target.ward=source.ward and target.scope_label=source.scope_label and target.building_type is not distinct from source.building_type and target.structure_type is not distinct from source.structure_type and target.survey_year=source.survey_year and target.observed_month is not distinct from source.observed_month) returning 1) select count(*) from added""")
        return {"inserted": int(inserted), "updated": int(updated), "skipped": 0}
    finally:
        await conn.close()


def _download(url: str, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    urllib.request.urlretrieve(url, path)
    return path


def parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", choices=("estate_housing_land", "estat_housing_land_122_5", "kouri"), default=None)
    parser.add_argument("--all", action="store_true")
    parser.add_argument("--file", type=Path, action="append", help="local XLSX; repeat for --all")
    parser.add_argument("--database-url", default=os.getenv("DATABASE_URL", ""))
    parser.add_argument("--dry-run", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    selected = ["estate_housing_land", "estat_housing_land_122_5", "kouri"] if args.all or not args.source else [args.source]
    files = list(args.file or [])
    rows: list[dict[str, Any]] = []
    total = {"skipped": 0, "unmapped_prefecture": 0, "unmapped_city": 0}
    for source in selected:
        url = {"estate_housing_land": HOUSING_URL, "estat_housing_land_122_5": HOUSING_122_5_URL, "kouri": KOURI_URL}[source]
        path = files.pop(0) if files else _download(url, Path("/tmp") / f"zouseeking-{source}.xlsx")
        parsed, report = (
            parse_housing_land_workbook(path) if source == "estate_housing_land"
            else parse_housing_land_122_5_workbook(path) if source == "estat_housing_land_122_5"
            else parse_kouri_workbook(path)
        )
        rows.extend(parsed)
        total["skipped"] += report.skipped
        total["unmapped_prefecture"] += report.unmapped_prefecture
        total["unmapped_city"] += report.unmapped_city
        print(f"source={source} parsed={len(parsed)} skipped={report.skipped} unmapped_prefecture={report.unmapped_prefecture} unmapped_city={report.unmapped_city}")
        if report.samples:
            print(f"unmapped_city_samples={','.join(report.samples[:5])}")
    if args.dry_run:
        result = {"inserted": 0, "updated": 0, "skipped": total["skipped"]}
    else:
        if not args.database_url:
            raise SystemExit("--database-url is required unless --dry-run is used")
        result = asyncio.run(upsert_rows(args.database_url, rows))
        result["skipped"] += total["skipped"]
    print("rent_reference_import " + " ".join(f"{key}={value}" for key, value in result.items()) + f" unmapped_prefecture={total['unmapped_prefecture']} unmapped_city={total['unmapped_city']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
