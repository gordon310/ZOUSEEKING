from pathlib import Path
import importlib.util
import sys

from openpyxl import Workbook

_SPEC = importlib.util.spec_from_file_location("import_rent_reference", Path(__file__).parents[2] / "scripts" / "import_rent_reference.py")
_MODULE = importlib.util.module_from_spec(_SPEC)
assert _SPEC.loader is not None
sys.modules[_SPEC.name] = _MODULE
_SPEC.loader.exec_module(_MODULE)
parse_housing_land_workbook = _MODULE.parse_housing_land_workbook
parse_kouri_workbook = _MODULE.parse_kouri_workbook


def _write_122_5_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "e122_5"
    sheet.append(["第１２２－５表 住宅の建て方(4区分)、構造(2区分)別"])
    sheet.append([None])
    sheet.append([None])
    sheet.append([None])
    sheet.append([None, None, None, "表章項目", None, "延べ面積１ｍ2当たり家賃"])
    sheet.append([None, None, None, "事項名", None, "住宅の家賃の平均"])
    sheet.append([None, None, None, "項目名", "00_総数", "2_家賃０円を含まない"])
    sheet.append([None, None, None, "表章単位", "戸", "円"])
    sheet.append(["地域識別コード", "地域区分", "住宅の建て方", "建物の構造", None, None])
    sheet.append(["a", "13000_東京都", "0_総数", "0_総数", 10, 1000])
    sheet.append(["a", "13000_東京都", "3_共同住宅", "2_非木造", 20, 2345])
    sheet.append(["a", "13000_東京都", "3_共同住宅", "2_非木造", 21, "-"])
    workbook.save(path)


def test_housing_land_parser_uses_excl_zero_and_maps_special_wards(tmp_path: Path):
    path = tmp_path / "housing.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "e122_4"
    sheet.append(["title"])
    sheet.append(["第122-4表"])
    for _ in range(2):
        sheet.append([])
    sheet.append([None, None, "表章項目", *([None] * 14)])
    sheet.append([None, None, "事項名", *([None] * 14)])
    sheet.append([None, None, "項目名", *([None] * 14)])
    sheet.append([None, None, "表章単位", *([None] * 14)])
    sheet.append(["地域識別コード", "地域区分", "住宅の所有の関係"] + [None] * 14)
    sheet.append(["a", "13000_東京都", "0_総数"] + [None] * 14)
    sheet.append(["0", "13101_千代田区", "3_民営借家"] + [None] * 12 + [3800, 3807])
    sheet.append(["1", "13100_特別区部", "3_民営借家"] + [None] * 12 + [2760, 2800])
    sheet.append(["2", "13102_中央区", "3_民営借家"] + [None] * 12 + ["-", "-"])
    workbook.save(path)

    rows, report = parse_housing_land_workbook(path)

    assert [(row["city"], row["ward"], row["rent_jpy_per_sqm_month_excl_zero"]) for row in rows] == [
        ("千代田区", "__not_subdivided__", 3807),
        ("东京23区", "__not_subdivided__", 2800),
    ]
    assert report.skipped == 1


def test_kouri_parser_converts_33sqm_month_to_sqm_month(tmp_path: Path):
    path = tmp_path / "kouri.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "B001-2"
    sheet.cell(10, 15, "地域")
    sheet.cell(10, 16, "01100")
    sheet.cell(11, 15, "地域")
    sheet.cell(11, 16, "札幌市")
    sheet.cell(12, 9, "2026年8月")
    sheet.cell(12, 10, 3001)
    sheet.cell(12, 11, "民営家賃")
    sheet.cell(12, 12, "1か月･3.3m2")
    sheet.cell(12, 16, 4233)
    workbook.save(path)

    rows, report = parse_kouri_workbook(path)

    assert rows[0]["city"] == "札幌市"
    assert rows[0]["observed_month"] == "2026-08"
    assert rows[0]["rent_jpy_per_sqm_month"] == 4233 / 3.3
    assert report.skipped == 0


def test_122_5_parser_uses_dimension_headers_and_skips_missing_rent(tmp_path):
    path = tmp_path / "122-5.xlsx"
    _write_122_5_fixture(path)

    rows, report = _MODULE.parse_housing_land_122_5_workbook(path)

    assert _MODULE.HOUSING_122_5_SOURCE == "estat_housing_land_122_5"
    assert report.skipped == 1
    assert len(rows) == 2
    assert rows[0]["building_type"] == "総数"
    assert rows[0]["structure_type"] == "総数"
    assert rows[1]["building_type"] == "共同住宅"
    assert rows[1]["structure_type"] == "非木造"
    assert rows[1]["rent_jpy_per_sqm_month_excl_zero"] == 2345
